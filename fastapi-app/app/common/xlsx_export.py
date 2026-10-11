# -*- coding: utf-8 -*-
"""xlsx 导出工具：统一构建与响应头（openpyxl 已在依赖中）。"""
from __future__ import annotations

import urllib.parse
from io import BytesIO

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_ROW_LIMIT = 50000        # 单次导出行数上限（超限截断）


def build_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet1")[:31]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = fill
    for r in rows[:EXPORT_ROW_LIMIT]:
        ws.append([v if v is not None else "" for v in r])
    # 列宽：按前 200 行内容长度粗略自适应（中文按 2 字符宽计）
    def _w(v) -> int:
        s = str(v)
        return sum(2 if ord(ch) > 127 else 1 for ch in s)

    for i, h in enumerate(headers, 1):
        width = max([_w(h)] + [_w(r[i - 1]) for r in rows[:200] if len(r) >= i] or [8])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 70)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def xlsx_response(filename: str, payload: bytes, truncated: bool = False) -> StreamingResponse:
    quoted = urllib.parse.quote(filename)
    return StreamingResponse(
        BytesIO(payload),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quoted}",
            "X-Row-Limit-Hit": "1" if truncated else "0",
        })
